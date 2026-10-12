"""
Shipyard AI export wrappers.

Accepts a list of shipyard record dictionaries and writes Excel or Word
output. If records contains a single item, the single-record layout is used;
otherwise the all-records/batch layout is used.
"""

import os
from typing import Dict, Any, List

from app.shipyard_ai.export_to_excel import (
    create_excel_with_all_shipyards,
    create_excel_from_template,
)
from app.shipyard_ai.export_to_word import create_word_document


LOGO_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "logo.png")


def _ensure_dir(output_path: str) -> None:
    directory = os.path.dirname(output_path)
    if directory:
        os.makedirs(directory, exist_ok=True)


def export_to_excel(records: List[Dict[str, Any]], output_path: str) -> str:
    """Export records to an Excel file and return the output filename."""
    _ensure_dir(output_path)

    if len(records) == 1:
        create_excel_from_template(records[0], output_path)
    else:
        create_excel_with_all_shipyards(records, output_path)

    return os.path.basename(output_path)


def export_to_word(records: List[Dict[str, Any]], output_path: str) -> str:
    """Export records to a Word file and return the output filename."""
    _ensure_dir(output_path)

    if len(records) == 1:
        create_word_document(records[0], output_path, logo_path=LOGO_PATH)
    else:
        # Batch layout: create a combined Word document with a section per record.
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn

        doc = Document()
        doc.styles["Normal"].font.name = "Arial"
        doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        doc.styles["Normal"].font.size = Pt(10.5)

        # Use header logo from the first section only (same logo for whole doc).
        from app.shipyard_ai.export_to_word import add_header_logo

        add_header_logo(doc, logo_path=LOGO_PATH)

        for idx, record in enumerate(records):
            if idx > 0:
                doc.add_page_break()

            title = doc.add_heading(
                record.get("English_name", record.get("shipyard_name", "Shipyard")),
                level=1,
            )
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if title.runs:
                run = title.runs[0]
                run.font.size = Pt(20)
                run.font.color.rgb = RGBColor(0, 51, 102)

            subtitle = doc.add_paragraph("Shipyard Profile & Introduction")
            subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if subtitle.runs:
                subtitle.runs[0].font.size = Pt(14)
                subtitle.runs[0].font.color.rgb = RGBColor(0, 51, 102)

            # Add a compact overview paragraph.
            overview = doc.add_paragraph()
            overview_text = (
                f"{record.get('English_name', 'The shipyard')} is located in "
                f"{record.get('location', 'unknown location')}."
            )
            if record.get("ubo"):
                overview_text += f" It is part of {record['ubo']} group."
            overview.add_run(overview_text)
            overview.paragraph_format.line_spacing = 1.5

            # Add basic info table.
            table = doc.add_table(rows=8, cols=2)
            table.style = "Light Grid Accent 1"
            rows = [
                ("Local Name", record.get("Local_name", "N/A")),
                ("Abbreviation", record.get("Abreviation", "N/A")),
                ("Location", record.get("location", "N/A")),
                ("Country", record.get("country", "N/A")),
                ("Establishment Date", record.get("establishment_date", "N/A")),
                ("Active Status", record.get("is_active", "N/A")),
                ("Shipyard Type", record.get("shipyard_type", "N/A")),
                ("UBO", record.get("ubo", "N/A")),
            ]
            for i, (label, value) in enumerate(rows):
                table.rows[i].cells[0].text = label
                table.rows[i].cells[1].text = str(value) if value is not None else "N/A"
                cell_run = table.rows[i].cells[0].paragraphs[0].runs[0]
                cell_run.font.bold = True

            doc.add_paragraph()

            # Production capacity.
            from app.shipyard_ai.export_to_word import _format_annual_capacity
            annual_vessel, annual_dwt = _format_annual_capacity(record.get("annual_production_capacity"))
            if annual_vessel or annual_dwt:
                p = doc.add_paragraph()
                p.add_run("Annual Production Capacity: ").bold = True
                parts = []
                if annual_vessel:
                    parts.append(f"Vessels: {annual_vessel}")
                if annual_dwt:
                    parts.append(f"DWT: {annual_dwt:,}")
                p.add_run("; ".join(parts))

            # Qualifications.
            certs = record.get("qualification_certifications")
            if certs:
                doc.add_heading("Qualifications and Certifications", level=2)
                for cert in str(certs).split(";"):
                    if cert.strip():
                        doc.add_paragraph(cert.strip(), style="List Bullet")

        doc.save(output_path)

    return os.path.basename(output_path)
